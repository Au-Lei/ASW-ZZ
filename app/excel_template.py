"""受控 Excel 模板版本和字段位置配置。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO

from app.extraction_schema import CORE_FIELD_NAMES


@dataclass(frozen=True, slots=True)
class CellLocation:
    sheet: str
    cell: str

    def __post_init__(self) -> None:
        if not self.sheet.strip() or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", self.cell):
            raise ValueError("工作表或单元格地址无效")


@dataclass(frozen=True, slots=True)
class TemplateMapping:
    version: str
    locations: tuple[tuple[str, CellLocation], ...]

    def __post_init__(self) -> None:
        if not self.version.strip():
            raise ValueError("模板版本不能为空")
        fields = tuple(name for name, _ in self.locations)
        if set(fields) != set(CORE_FIELD_NAMES) or len(fields) != len(CORE_FIELD_NAMES):
            raise ValueError("模板映射必须包含全部核心字段且不得重复")
        targets = tuple(location for _, location in self.locations)
        if len(targets) != len(set(targets)):
            raise ValueError("模板映射不能让多个字段写入同一单元格")


def validate_template_mapping(template_bytes: bytes, mapping: TemplateMapping) -> None:
    """验证映射的工作表和目标格，禁止覆盖公式及合并区域内部格。"""

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(template_bytes), read_only=False, data_only=False, keep_links=False)
        for field_name, location in mapping.locations:
            if location.sheet not in workbook.sheetnames:
                raise ValueError(f"字段 {field_name} 的目标工作表不存在")
            sheet = workbook[location.sheet]
            cell = sheet[location.cell]
            if cell.data_type == "f":
                raise ValueError(f"字段 {field_name} 的目标格包含公式")
            if cell.__class__.__name__ == "MergedCell":
                raise ValueError(f"字段 {field_name} 的目标格位于合并区域内部")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel 模板无法打开或映射无效") from None

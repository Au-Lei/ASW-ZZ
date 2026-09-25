"""模板映射配置与最小 Excel 引擎验证。"""

import io
import unittest

from app.excel_template import CellLocation, TemplateMapping, validate_template_mapping
from app.extraction_schema import CORE_FIELD_NAMES


def _mapping() -> TemplateMapping:
    return TemplateMapping("template-v1", tuple(
        (name, CellLocation("联系单", f"B{index}"))
        for index, name in enumerate(CORE_FIELD_NAMES, start=2)
    ))


class TemplateMappingTests(unittest.TestCase):
    def test_requires_complete_unique_versioned_mapping(self) -> None:
        mapping = _mapping()
        self.assertEqual(len(mapping.locations), len(CORE_FIELD_NAMES))
        with self.assertRaises(ValueError):
            TemplateMapping(" ", mapping.locations)
        with self.assertRaises(ValueError):
            TemplateMapping("v1", mapping.locations[:-1])
        with self.assertRaises(ValueError):
            CellLocation("联系单", "B0")

    def test_opens_saves_and_validates_synthetic_template(self) -> None:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "联系单"
        sheet["A1"] = "业务联系单"
        sheet["A1"].font = Font(bold=True)
        sheet["D2"] = "=1+1"
        sheet.merge_cells("F2:G2")
        sheet.print_area = "A1:G20"
        data = io.BytesIO()
        workbook.save(data)
        validate_template_mapping(data.getvalue(), _mapping())
        reopened = load_workbook(io.BytesIO(data.getvalue()), data_only=False)
        self.assertTrue(reopened["联系单"]["A1"].font.bold)
        self.assertEqual(reopened["联系单"]["D2"].value, "=1+1")
        self.assertIn("F2:G2", {str(item) for item in reopened["联系单"].merged_cells.ranges})
        self.assertIn("$A$1:$G$20", reopened["联系单"].print_area)

        invalid = TemplateMapping("v2", tuple(
            (name, CellLocation("联系单", "D2" if index == 0 else f"B{index+2}"))
            for index, name in enumerate(CORE_FIELD_NAMES)
        ))
        with self.assertRaises(ValueError):
            validate_template_mapping(data.getvalue(), invalid)


if __name__ == "__main__":
    unittest.main()
